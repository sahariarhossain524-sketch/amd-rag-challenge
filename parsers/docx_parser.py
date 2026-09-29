import docx

def parse_docx(file_path, rel_path):
    """
    Parses DOCX documents.
    Extracts text from paragraphs and embedded tables.
    """
    chunks = []
    try:
        doc = docx.Document(file_path)
        full_text = []

        # 1. Paragraphs
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                full_text.append(t)

        # 2. Tables (crucial for roadmap quarters, part tables, etc.)
        for t_idx, table in enumerate(doc.tables):
            table_lines = []
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                # remove adjacent duplicates caused by merged cells
                dedup_cells = []
                for c in cells:
                    if not dedup_cells or c != dedup_cells[-1]:
                        dedup_cells.append(c)
                if dedup_cells:
                    table_lines.append(" | ".join(dedup_cells))
            if table_lines:
                full_text.append(f"\n--- Table {t_idx + 1} ---\n" + "\n".join(table_lines))

        content = "\n".join(full_text).strip()
        if content:
            chunks.append({
                "file_path": rel_path,
                "source_type": "docx",
                "text": content
            })
    except Exception:
        return []
    
    return chunks
