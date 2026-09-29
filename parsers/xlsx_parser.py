import openpyxl

def parse_xlsx(file_path, rel_path):
    """
    Parses multi-sheet XLSX spreadsheets.
    Formats rows with column headers for semantic and exact retrieval.
    """
    chunks = []
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            # First non-empty row as header
            header_idx = -1
            headers = []
            for idx, r in enumerate(rows):
                if any(cell is not None for cell in r):
                    header_idx = idx
                    headers = [str(c).strip() if c is not None else f"Col_{i}" for i, c in enumerate(r)]
                    break

            if header_idx == -1:
                continue

            sheet_lines = [f"Sheet: {sheet_name}"]
            for r in rows[header_idx + 1:]:
                if not any(cell is not None for cell in r):
                    continue
                row_items = []
                for h, val in zip(headers, r):
                    if val is not None and str(val).strip():
                        row_items.append(f"{h}: {str(val).strip()}")
                if row_items:
                    sheet_lines.append(" | ".join(row_items))

            sheet_text = "\n".join(sheet_lines).strip()
            if sheet_text:
                chunks.append({
                    "file_path": rel_path,
                    "source_type": "xlsx",
                    "sheet": sheet_name,
                    "text": sheet_text
                })
        wb.close()
    except Exception:
        return []

    return chunks
